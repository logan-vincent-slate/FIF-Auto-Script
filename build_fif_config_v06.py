#!/usr/bin/env python3
"""
build_fif_config_gui.py  (v05 - GUI)

Window-based builder for the two files an IPA "FindInFiles /
Expression-Based Events" script needs:

    1. A .asl signal-lookup / config file (JSON)
    2. A matching Excel workbook (.xlsx), one worksheet per event

Run it (no arguments) and a window opens:

    python build_fif_config_gui.py

Tabs:
    1. Files    - output folder, file name, DBC folder (all DBCs load automatically)
    2. Signals  - search the DBCs, add signals; the FIRST one drives the time basis
    3. Events   - name / StartExpression / EndExpression for each event
Then press "Generate files".

Max(), Min() and Avg() are automatically added as EventBasedMetrics for
EVERY tracked signal on EVERY event.

Requires:  pip install openpyxl cantools
"""

import glob
import json
import os
import re
import sys
import tkinter as tk
from tkinter import ttk, filedialog, messagebox, simpledialog, scrolledtext

try:
    import openpyxl
    from openpyxl.styles import Font, Alignment, Border, Side
except ImportError:
    openpyxl = None

try:
    import cantools
except ImportError:
    cantools = None

# Default folder holding the DBC files.
DEFAULT_DBC_DIR = r"C:\Users\LoganVincent\OneDrive - SLATE\Documents\CMS_6.5_DBC"

# Remembers which network label (e.g. "BODY CAN") belongs to each DBC file.
NETWORK_MAP_FILE = os.path.join(
    os.path.dirname(os.path.abspath(__file__)), "dbc_network_map.json"
)

# Remembers the last output folder you used.
SETTINGS_FILE = os.path.join(
    os.path.dirname(os.path.abspath(__file__)), "fif_builder_settings.json"
)

MAX_RESULTS_SHOWN = 300

METRIC_TYPES = ["Max", "Min", "Avg"]

BASE_COLUMNS = [
    "EventIndex",
    "Vehicle",
    "FileID",
    "EventFilename",
    "EventStartTime",
    "EventDuration",
]

INVALID_SHEET_CHARS = re.compile(r"[\[\]\:\*\?/\\]")

def load_settings():
    try:
        with open(SETTINGS_FILE) as f:
            return json.load(f)
    except (OSError, ValueError):
        return {}


def save_settings(settings):
    try:
        with open(SETTINGS_FILE, "w") as f:
            json.dump(settings, f, indent=4)
    except OSError:
        pass


def folder_is_writable(folder):
    """True if we can create a file in `folder` (creating the folder if needed)."""
    try:
        os.makedirs(folder, exist_ok=True)
        test = os.path.join(folder, ".fif_write_test")
        with open(test, "w") as f:
            f.write("x")
        os.remove(test)
        return True
    except OSError:
        return False


def default_output_dir():
    """Last-used folder if it still works, else Documents, else home."""
    last = load_settings().get("output_dir", "")
    if last and os.path.isdir(last) and folder_is_writable(last):
        return last
    docs = os.path.join(os.path.expanduser("~"), "Documents")
    if os.path.isdir(docs) and folder_is_writable(docs):
        return docs
    return os.path.expanduser("~")


def make_signal_entry(name_in_script, channel_name, message_name, network_name):
    return {
        "name_in_script": name_in_script,
        "OptionalList": [
            {
                "channel_name": channel_name,
                "message_name": message_name,
                "network_name": network_name,
            }
        ],
    }


def load_network_map():
    try:
        with open(NETWORK_MAP_FILE) as f:
            return json.load(f)
    except (OSError, ValueError):
        return {}


def save_network_map(mapping):
    try:
        with open(NETWORK_MAP_FILE, "w") as f:
            json.dump(mapping, f, indent=4)
    except OSError:
        print("  (could not save network map - you'll be asked again next time)")


def find_dbc_files(folder):
    files = glob.glob(os.path.join(folder, "*.dbc")) + glob.glob(
        os.path.join(folder, "*.DBC")
    )
    return sorted(set(files))


def search_items(items, query):
    terms = query.lower().split()
    out = []
    for it in items:
        hay = f"{it['signal']} {it['message']}".lower()
        if all(t in hay for t in terms):
            out.append(it)
    return out


def auto_metrics_for_event(signal_names):
    """Max, Min and Avg for every tracked signal - grouped by signal."""
    metrics = []
    for sig in signal_names:
        for metric_type in METRIC_TYPES:
            metrics.append(f"{metric_type}({sig})")
    return metrics


def build_asl_dict(signals, time_basis, events, template_filename):
    return {
        "SignalListForUseInTimeBasis": time_basis,
        "DefaultDataSpyConfigFile": "",
        "TemplateFilename": template_filename,
        "Channels": signals,
        "EventDefinitions": events,
    }


def write_asl_file(path, asl_dict):
    with open(path, "w") as f:
        json.dump(asl_dict, f, indent=4)


def safe_sheet_name(name, used_names):
    clean = INVALID_SHEET_CHARS.sub("", name).strip() or "Event"
    clean = clean[:31]
    candidate = clean
    counter = 2
    while candidate.lower() in used_names:
        suffix = f" ({counter})"
        candidate = clean[: 31 - len(suffix)] + suffix
        counter += 1
    used_names.add(candidate.lower())
    return candidate


def build_excel(path, events):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    header_font = Font(bold=True)
    header_align = Alignment(horizontal="left", wrap_text=True)
    top_border = Side(style="medium")
    bottom_border = Side(style="thick")
    header_border = Border(top=top_border, bottom=bottom_border)

    used_names = set()
    for event in events:
        sheet_name = safe_sheet_name(event["Description"], used_names)
        ws = wb.create_sheet(title=sheet_name)

        headers = BASE_COLUMNS + list(event["EventBasedMetrics"])

        # Column B is the first data column (column A is left as a spacer,
        # matching the reference template).
        ws.column_dimensions["A"].width = 3.44

        for i, header in enumerate(headers):
            col = 2 + i  # start at column B
            cell = ws.cell(row=2, column=col, value=header)
            cell.font = header_font
            cell.alignment = header_align
            cell.border = header_border

            col_letter = openpyxl.utils.get_column_letter(col)
            if header == "EventFilename":
                width = 44.11
            elif header == "EventStartTime":
                width = 18.44
            elif header == "EventDuration":
                width = 13.44
            elif header == "EventIndex":
                width = 10.78
            elif header in ("Vehicle", "FileID"):
                width = 12.0
            else:
                # metric columns - size to fit the header text
                width = max(18.0, len(header) + 6)
            ws.column_dimensions[col_letter].width = width

        ws.row_dimensions[2].height = 15.0

        # Freeze the header row and the base (non-metric) columns.
        freeze_col_letter = openpyxl.utils.get_column_letter(2 + len(BASE_COLUMNS))
        ws.freeze_panes = f"{freeze_col_letter}4"

    wb.save(path)
# ----------------------------------------------------------------------
# GUI
# ----------------------------------------------------------------------

class App(tk.Tk):
    def __init__(self):
        super().__init__()
        self.title("FIF Event Config + Excel Template Builder")
        self.geometry("1020x720")
        self.minsize(900, 620)

        self.items = []        # every signal found in the DBCs
        self.net_vars = {}     # dbc filename -> StringVar (network name)
        self.results = []      # current search results
        self.picked = []       # chosen signals (first = time basis)
        self.events = []       # defined events
        self._last_expr_entry = None

        self._build_ui()
        self.after(150, self.load_dbcs)

    # ------------------------------------------------------------ layout
    def _build_ui(self):
        pad = {"padx": 8, "pady": 4}

        nb = ttk.Notebook(self)
        nb.pack(fill="both", expand=True, padx=8, pady=(8, 0))
        self.nb = nb

        self.tab_files = ttk.Frame(nb)
        self.tab_sigs = ttk.Frame(nb)
        self.tab_events = ttk.Frame(nb)
        nb.add(self.tab_files, text="  1. Files  ")
        nb.add(self.tab_sigs, text="  2. Signals  ")
        nb.add(self.tab_events, text="  3. Events  ")

        self._build_files_tab(pad)
        self._build_signals_tab(pad)
        self._build_events_tab(pad)

        bar = ttk.Frame(self)
        bar.pack(fill="x", padx=8, pady=8)
        self.summary_var = tk.StringVar()
        ttk.Label(bar, textvariable=self.summary_var).pack(side="left")
        ttk.Button(bar, text="Generate files", command=self.generate).pack(side="right")
        ttk.Button(bar, text="Preview config", command=self.preview).pack(side="right", padx=6)
        self._update_summary()

    # ---- tab 1
    def _build_files_tab(self, pad):
        f = self.tab_files
        f.columnconfigure(1, weight=1)

        ttk.Label(f, text="Output folder", font=("", 10, "bold")).grid(row=0, column=0, sticky="w", **pad)
        self.out_dir = tk.StringVar(value=default_output_dir())
        ttk.Entry(f, textvariable=self.out_dir).grid(row=0, column=1, sticky="ew", **pad)
        ttk.Button(f, text="Browse...", command=self._browse_out).grid(row=0, column=2, **pad)
        ttk.Label(f, text="Where the .asl and .xlsx files will be saved (remembered for next time).",
                  foreground="#666").grid(row=1, column=1, sticky="w", padx=8)

        ttk.Label(f, text="Base file name", font=("", 10, "bold")).grid(row=2, column=0, sticky="w", **pad)
        self.base_name = tk.StringVar(value="EventConfig")
        ttk.Entry(f, textvariable=self.base_name).grid(row=2, column=1, sticky="ew", **pad)
        ttk.Label(f, text="No extension. Example: Config_BrakeTest creates "
                          "Config_BrakeTest.asl and Config_BrakeTest.xlsx",
                  foreground="#666", wraplength=620, justify="left").grid(row=3, column=1, sticky="w", padx=8)

        ttk.Separator(f).grid(row=4, column=0, columnspan=3, sticky="ew", pady=12)

        ttk.Label(f, text="DBC folder", font=("", 10, "bold")).grid(row=5, column=0, sticky="w", **pad)
        self.dbc_dir = tk.StringVar(value=DEFAULT_DBC_DIR)
        ttk.Entry(f, textvariable=self.dbc_dir).grid(row=5, column=1, sticky="ew", **pad)
        btns = ttk.Frame(f)
        btns.grid(row=5, column=2, **pad)
        ttk.Button(btns, text="Browse...", command=self._browse_dbc).pack(side="left")
        ttk.Button(btns, text="Reload", command=self.load_dbcs).pack(side="left", padx=(4, 0))

        self.dbc_status = tk.StringVar(value="")
        ttk.Label(f, textvariable=self.dbc_status, foreground="#666").grid(
            row=6, column=1, sticky="w", padx=8)

        self.net_frame = ttk.LabelFrame(
            f, text="Network name for each DBC (optional - edit if the default is wrong)")
        self.net_frame.grid(row=7, column=0, columnspan=3, sticky="ew", padx=8, pady=14)
        self.net_frame.columnconfigure(1, weight=1)
        ttk.Label(
            f,
            text=("Network name = the CAN bus label used in your data files, e.g. "
                  "\"BODY CAN\" or \"HV CAN\". It is written into network_name for "
                  "every signal picked from that DBC and remembered for next time."),
            wraplength=900, foreground="#666",
        ).grid(row=8, column=0, columnspan=3, sticky="w", padx=8)

    # ---- tab 2
    def _build_signals_tab(self, pad):
        f = self.tab_sigs
        f.columnconfigure(0, weight=3)
        f.columnconfigure(1, weight=2)
        f.rowconfigure(2, weight=1)

        ttk.Label(
            f,
            text=("Type part of a signal or message name (e.g. \"brake\"). Select one or more "
                  "results and click Add (or double-click). The FIRST signal in your list drives "
                  "the time basis, so pick your most important / most frequently sent signal first."),
            wraplength=960, foreground="#666",
        ).grid(row=0, column=0, columnspan=2, sticky="w", **pad)

        # left: search
        left = ttk.LabelFrame(f, text="Search DBC signals")
        left.grid(row=1, column=0, rowspan=2, sticky="nsew", padx=8, pady=4)
        left.columnconfigure(0, weight=1)
        left.rowconfigure(2, weight=1)
        self.search_var = tk.StringVar()
        self.search_var.trace_add("write", lambda *_: self.do_search())
        e = ttk.Entry(left, textvariable=self.search_var)
        e.grid(row=0, column=0, sticky="ew", padx=6, pady=6)
        self.search_entry = e
        self.search_info = tk.StringVar(value="")
        ttk.Label(left, textvariable=self.search_info, foreground="#666").grid(
            row=1, column=0, sticky="w", padx=6)
        wrap = ttk.Frame(left)
        wrap.grid(row=2, column=0, sticky="nsew", padx=6, pady=6)
        wrap.rowconfigure(0, weight=1)
        wrap.columnconfigure(0, weight=1)
        self.res_list = tk.Listbox(wrap, selectmode="extended", activestyle="none", exportselection=False)
        sb = ttk.Scrollbar(wrap, command=self.res_list.yview)
        sbx = ttk.Scrollbar(wrap, orient="horizontal", command=self.res_list.xview)
        self.res_list.configure(yscrollcommand=sb.set, xscrollcommand=sbx.set)
        self.res_list.grid(row=0, column=0, sticky="nsew")
        sb.grid(row=0, column=1, sticky="ns")
        sbx.grid(row=1, column=0, sticky="ew")
        self.res_list.bind("<Double-Button-1>", lambda e: self.add_selected())
        ttk.Button(left, text="Add selected  >>", command=self.add_selected).grid(
            row=3, column=0, sticky="e", padx=6, pady=6)

        # right: picked
        right = ttk.LabelFrame(f, text="Signals to track (first = time basis)")
        right.grid(row=1, column=1, rowspan=2, sticky="nsew", padx=8, pady=4)
        right.columnconfigure(0, weight=1)
        right.rowconfigure(0, weight=1)
        wrap2 = ttk.Frame(right)
        wrap2.grid(row=0, column=0, sticky="nsew", padx=6, pady=6)
        wrap2.rowconfigure(0, weight=1)
        wrap2.columnconfigure(0, weight=1)
        self.pick_list = tk.Listbox(wrap2, selectmode="extended", activestyle="none", exportselection=False)
        sb2 = ttk.Scrollbar(wrap2, command=self.pick_list.yview)
        self.pick_list.configure(yscrollcommand=sb2.set)
        self.pick_list.grid(row=0, column=0, sticky="nsew")
        sb2.grid(row=0, column=1, sticky="ns")
        pb = ttk.Frame(right)
        pb.grid(row=1, column=0, sticky="ew", padx=6, pady=6)
        ttk.Button(pb, text="Remove", command=self.remove_picked).pack(side="left")
        ttk.Button(pb, text="Move up", command=lambda: self.move_picked(-1)).pack(side="left", padx=4)
        ttk.Button(pb, text="Move down", command=lambda: self.move_picked(1)).pack(side="left")
        ttk.Button(pb, text="Add manually...", command=self.add_manual).pack(side="right")

    # ---- tab 3
    def _build_events_tab(self, pad):
        f = self.tab_events
        f.columnconfigure(0, weight=1)
        f.columnconfigure(1, weight=2)
        f.rowconfigure(1, weight=1)

        ttk.Label(
            f,
            text=("Each event is one start/end condition and gets its own Excel sheet. Fill in the "
                  "form on the right, then click \"Add event\". Max, Min and Avg are captured "
                  "automatically for every tracked signal."),
            wraplength=960, foreground="#666",
        ).grid(row=0, column=0, columnspan=2, sticky="w", **pad)

        # left: event list
        left = ttk.LabelFrame(f, text="Events")
        left.grid(row=1, column=0, sticky="nsew", padx=8, pady=4)
        left.columnconfigure(0, weight=1)
        left.rowconfigure(0, weight=1)
        wrap = ttk.Frame(left)
        wrap.grid(row=0, column=0, sticky="nsew", padx=6, pady=6)
        wrap.rowconfigure(0, weight=1)
        wrap.columnconfigure(0, weight=1)
        self.ev_list = tk.Listbox(wrap, activestyle="none", exportselection=False)
        sb = ttk.Scrollbar(wrap, command=self.ev_list.yview)
        self.ev_list.configure(yscrollcommand=sb.set)
        self.ev_list.grid(row=0, column=0, sticky="nsew")
        sb.grid(row=0, column=1, sticky="ns")
        self.ev_list.bind("<<ListboxSelect>>", self.on_event_select)
        eb = ttk.Frame(left)
        eb.grid(row=1, column=0, sticky="ew", padx=6, pady=6)
        ttk.Button(eb, text="Delete selected", command=self.delete_event).pack(side="left")
        ttk.Button(eb, text="New (clear form)", command=self.clear_event_form).pack(side="right")

        # right: form
        right = ttk.LabelFrame(f, text="Event details")
        right.grid(row=1, column=1, sticky="nsew", padx=8, pady=4)
        right.columnconfigure(0, weight=1)

        ttk.Label(right, text="Name of event", font=("", 10, "bold")).grid(row=0, column=0, sticky="w", padx=8, pady=(8, 0))
        ttk.Label(right, text="Short label. Becomes the Description in the config AND the Excel "
                              "sheet name (max 31 characters).",
                  foreground="#666", wraplength=560).grid(row=1, column=0, sticky="w", padx=8)
        self.ev_name = tk.StringVar()
        ttk.Entry(right, textvariable=self.ev_name).grid(row=2, column=0, sticky="ew", padx=8, pady=(2, 0))
        self.ex_name = tk.StringVar(value="example:  Brake Pedal Pressed")
        ttk.Label(right, textvariable=self.ex_name, foreground="#2a6").grid(row=3, column=0, sticky="w", padx=8)

        ttk.Label(right, text="StartExpression", font=("", 10, "bold")).grid(row=4, column=0, sticky="w", padx=8, pady=(10, 0))
        ttk.Label(right, text="The condition that STARTS the event. Use signal names exactly as "
                              "listed below and wrap the condition in parentheses.",
                  foreground="#666", wraplength=560).grid(row=5, column=0, sticky="w", padx=8)
        self.ev_start = tk.StringVar()
        self.start_entry = ttk.Entry(right, textvariable=self.ev_start)
        self.start_entry.grid(row=6, column=0, sticky="ew", padx=8, pady=(2, 0))
        self.start_entry.bind("<FocusIn>", lambda e: setattr(self, "_last_expr_entry", self.start_entry))
        self._last_expr_entry = self.start_entry
        self.ex_start = tk.StringVar()
        ttk.Label(right, textvariable=self.ex_start, foreground="#2a6").grid(row=7, column=0, sticky="w", padx=8)

        ttk.Label(right, text="EndExpression", font=("", 10, "bold")).grid(row=8, column=0, sticky="w", padx=8, pady=(10, 0))
        ttk.Label(right, text="The condition that ENDS the event. For analog signals use two "
                              "different thresholds so the event doesn't flicker, e.g. start "
                              "(Voltage < 13.2), end (Voltage > 13.3).",
                  foreground="#666", wraplength=560).grid(row=9, column=0, sticky="w", padx=8)
        self.ev_end = tk.StringVar()
        self.end_entry = ttk.Entry(right, textvariable=self.ev_end)
        self.end_entry.grid(row=10, column=0, sticky="ew", padx=8, pady=(2, 0))
        self.end_entry.bind("<FocusIn>", lambda e: setattr(self, "_last_expr_entry", self.end_entry))
        self.ex_end = tk.StringVar()
        ttk.Label(right, textvariable=self.ex_end, foreground="#2a6").grid(row=11, column=0, sticky="w", padx=8)

        ttk.Label(right, text="Signals you can use (double-click to insert into the expression you're editing):",
                  foreground="#666").grid(row=12, column=0, sticky="w", padx=8, pady=(10, 0))
        self.sig_help = tk.Listbox(right, height=5, activestyle="none", exportselection=False)
        self.sig_help.grid(row=13, column=0, sticky="ew", padx=8, pady=2)
        self.sig_help.bind("<Double-Button-1>", self.insert_signal_name)

        bb = ttk.Frame(right)
        bb.grid(row=14, column=0, sticky="e", padx=8, pady=10)
        ttk.Button(bb, text="Update selected event", command=self.update_event).pack(side="left", padx=4)
        ttk.Button(bb, text="Add event", command=self.add_event).pack(side="left")
        self._refresh_examples()

    # ------------------------------------------------------------ helpers
    def _browse_out(self):
        d = filedialog.askdirectory(initialdir=self.out_dir.get() or os.getcwd(), title="Output folder")
        if d:
            self.out_dir.set(d)

    def _browse_dbc(self):
        d = filedialog.askdirectory(initialdir=self.dbc_dir.get() or os.getcwd(), title="DBC folder")
        if d:
            self.dbc_dir.set(d)
            self.load_dbcs()

    def _update_summary(self):
        self.summary_var.set(f"{len(self.picked)} signal(s)  |  {len(self.events)} event(s)")

    def _network_for(self, dbc_key):
        var = self.net_vars.get(dbc_key)
        return var.get().strip() if var else ""

    def _refresh_examples(self):
        sig = self.picked[0]["name_in_script"] if self.picked else "YourSignalName"
        self.ex_start.set(f"example:  ({sig} > 0)")
        self.ex_end.set(f"example:  ({sig} < 1)")
        self.sig_help.delete(0, "end")
        for p in self.picked:
            self.sig_help.insert("end", p["name_in_script"])

    # ------------------------------------------------------------ DBC loading
    def load_dbcs(self):
        if cantools is None:
            self.dbc_status.set("cantools is not installed (pip install cantools). "
                                "You can still add signals manually on the Signals tab.")
            return
        folder = self.dbc_dir.get().strip().strip('"')
        files = find_dbc_files(folder) if folder else []
        for w in self.net_frame.winfo_children():
            w.destroy()
        self.items = []
        self.net_vars = {}
        if not files:
            self.dbc_status.set(f"No .dbc files found in: {folder or '(no folder set)'}  "
                                "- use Browse... to pick the folder.")
            self.do_search()
            return

        self.dbc_status.set(f"Loading {len(files)} DBC file(s)...")
        self.update_idletasks()
        net_map = load_network_map()
        loaded = 0
        errors = []
        for row, path in enumerate(files):
            key = os.path.basename(path)
            var = tk.StringVar(value=net_map.get(key, os.path.splitext(key)[0]))
            try:
                db = cantools.database.load_file(path, strict=False)
            except Exception as e:  # noqa: BLE001
                errors.append(f"{key}: {e}")
                continue
            count = 0
            for msg in db.messages:
                for sig in msg.signals:
                    self.items.append({"signal": sig.name, "message": msg.name, "dbc": key})
                    count += 1
            self.net_vars[key] = var
            ttk.Label(self.net_frame, text=f"{key}  ({count} signals)").grid(
                row=row, column=0, sticky="w", padx=8, pady=3)
            ttk.Entry(self.net_frame, textvariable=var).grid(
                row=row, column=1, sticky="ew", padx=8, pady=3)
            loaded += 1
        msg = f"Loaded {loaded} of {len(files)} DBC file(s) from {folder}  ({len(self.items)} signals)"
        if errors:
            msg += "\nFailed: " + "; ".join(errors)
        self.dbc_status.set(msg)
        self.do_search()

    # ------------------------------------------------------------ signal search/pick
    def do_search(self):
        q = self.search_var.get().strip()
        self.res_list.delete(0, "end")
        self.results = []
        if not q:
            self.search_info.set("Type to search..." if self.items else "No DBC signals loaded.")
            return
        terms = q.lower().split()
        matches = [it for it in self.items
                   if all(t in f"{it['signal']} {it['message']}".lower() for t in terms)]
        self.results = matches[:MAX_RESULTS_SHOWN]
        for it in self.results:
            self.res_list.insert("end", f"[{self._network_for(it['dbc'])}]  {it['signal']}    ({it['message']})")
        extra = len(matches) - len(self.results)
        self.search_info.set(f"{len(matches)} match(es)" + (f" - showing first {len(self.results)}, refine to narrow" if extra > 0 else ""))

    def _unique_name(self, name):
        existing = {p["name_in_script"] for p in self.picked}
        while name in existing or not name:
            alt = simpledialog.askstring(
                "Name already used",
                f"'{name}' is already used as a name_in_script.\nEnter a different name for this signal:",
                parent=self)
            if alt is None:
                return None
            name = alt.strip()
        return name

    def add_selected(self):
        sel = self.res_list.curselection()
        if not sel:
            return
        for idx in sel:
            it = self.results[idx]
            if any(p.get("dbc") == it["dbc"] and p["signal"] == it["signal"] and p["message"] == it["message"]
                   for p in self.picked):
                continue
            name = self._unique_name(it["signal"])
            if name is None:
                continue
            self.picked.append({"name_in_script": name, "signal": it["signal"],
                                "message": it["message"], "dbc": it["dbc"], "network": None})
        self._refresh_picked()

    def add_manual(self):
        name = simpledialog.askstring("Add manually", "name_in_script (used inside expressions):", parent=self)
        if not name or not name.strip():
            return
        msg = simpledialog.askstring("Add manually", "message_name:", parent=self)
        if not msg or not msg.strip():
            return
        net = simpledialog.askstring("Add manually", "network_name (e.g. BODY CAN):", parent=self) or ""
        name = self._unique_name(name.strip())
        if name is None:
            return
        self.picked.append({"name_in_script": name, "signal": name, "message": msg.strip(),
                            "dbc": None, "network": net.strip()})
        self._refresh_picked()

    def remove_picked(self):
        for idx in reversed(self.pick_list.curselection()):
            del self.picked[idx]
        self._refresh_picked()

    def move_picked(self, delta):
        sel = list(self.pick_list.curselection())
        if len(sel) != 1:
            return
        i = sel[0]
        j = i + delta
        if 0 <= j < len(self.picked):
            self.picked[i], self.picked[j] = self.picked[j], self.picked[i]
            self._refresh_picked()
            self.pick_list.selection_set(j)

    def _picked_network(self, p):
        return p["network"] if p["dbc"] is None else self._network_for(p["dbc"])

    def _refresh_picked(self):
        self.pick_list.delete(0, "end")
        for i, p in enumerate(self.picked):
            tag = "     <- time basis" if i == 0 else ""
            self.pick_list.insert(
                "end", f"{i + 1}. {p['name_in_script']}    [{p['message']} | {self._picked_network(p)}]{tag}")
        self._refresh_examples()
        self._update_summary()

    # ------------------------------------------------------------ events
    def insert_signal_name(self, _evt=None):
        sel = self.sig_help.curselection()
        if not sel:
            return
        name = self.sig_help.get(sel[0])
        entry = self._last_expr_entry or self.start_entry
        entry.insert(entry.index("insert"), name)
        entry.focus_set()

    def _read_form(self):
        name = self.ev_name.get().strip()
        start = self.ev_start.get().strip()
        end = self.ev_end.get().strip()
        if not name or not start or not end:
            messagebox.showwarning("Missing information",
                                   "Please fill in the name, StartExpression and EndExpression.", parent=self)
            return None
        names = [p["name_in_script"] for p in self.picked]
        if names and not any(n in start or n in end for n in names):
            if not messagebox.askyesno(
                    "Check expressions",
                    "Neither expression mentions any of your tracked signals.\n\nSave it anyway?",
                    parent=self):
                return None
        return {"Description": name, "StartExpression": start, "EndExpression": end}

    def _refresh_events(self):
        self.ev_list.delete(0, "end")
        for e in self.events:
            self.ev_list.insert("end", f"{e['Description']}   :   {e['StartExpression']}  ->  {e['EndExpression']}")
        self._update_summary()

    def add_event(self):
        ev = self._read_form()
        if ev:
            self.events.append(ev)
            self._refresh_events()
            self.clear_event_form()

    def update_event(self):
        sel = self.ev_list.curselection()
        if not sel:
            messagebox.showinfo("Update event", "Select an event in the list first.", parent=self)
            return
        ev = self._read_form()
        if ev:
            self.events[sel[0]] = ev
            self._refresh_events()
            self.ev_list.selection_set(sel[0])

    def delete_event(self):
        sel = self.ev_list.curselection()
        if sel:
            del self.events[sel[0]]
            self._refresh_events()
            self.clear_event_form()

    def clear_event_form(self):
        self.ev_name.set("")
        self.ev_start.set("")
        self.ev_end.set("")
        self.ev_list.selection_clear(0, "end")

    def on_event_select(self, _evt=None):
        sel = self.ev_list.curselection()
        if sel:
            e = self.events[sel[0]]
            self.ev_name.set(e["Description"])
            self.ev_start.set(e["StartExpression"])
            self.ev_end.set(e["EndExpression"])

    # ------------------------------------------------------------ output
    def _build_config(self):
        signals = [make_signal_entry(p["name_in_script"], p["signal"], p["message"], self._picked_network(p))
                   for p in self.picked]
        time_basis = [signals[0]["name_in_script"]]
        names = [s["name_in_script"] for s in signals]
        events = [dict(e, EventBasedMetrics=auto_metrics_for_event(names)) for e in self.events]
        base = self.base_name.get().strip()
        return signals, events, build_asl_dict(signals, time_basis, events, f"{base}.xlsx")

    def _validate(self):
        if not self.base_name.get().strip():
            self.nb.select(self.tab_files)
            messagebox.showwarning("Missing information", "Enter a base file name.", parent=self)
            return False
        if not self.picked:
            self.nb.select(self.tab_sigs)
            messagebox.showwarning("Missing information", "Add at least one signal.", parent=self)
            return False
        if not self.events:
            self.nb.select(self.tab_events)
            messagebox.showwarning("Missing information",
                                   "Add at least one event (fill in the form and click \"Add event\").", parent=self)
            return False
        return True

    def preview(self):
        if not self._validate():
            return
        _, _, cfg = self._build_config()
        win = tk.Toplevel(self)
        win.title("Config preview (.asl)")
        win.geometry("700x600")
        txt = scrolledtext.ScrolledText(win, font=("Consolas", 10))
        txt.pack(fill="both", expand=True)
        txt.insert("1.0", json.dumps(cfg, indent=4))
        txt.configure(state="disabled")

    def generate(self):
        if not self._validate():
            return
        if openpyxl is None:
            messagebox.showerror("Missing package", "openpyxl is required:\n    pip install openpyxl", parent=self)
            return
        out_dir = self.out_dir.get().strip().strip('"')
        base = self.base_name.get().strip()
        if not out_dir or not folder_is_writable(out_dir):
            self.nb.select(self.tab_files)
            messagebox.showwarning(
                "Can't save there",
                f"The program can't write to this folder:\n{out_dir or '(none set)'}\n\n"
                "Choose a different output folder, such as Documents or a folder on your Desktop.",
                parent=self)
            return
        asl_path = os.path.join(out_dir, f"{base}.asl")
        xlsx_path = os.path.join(out_dir, f"{base}.xlsx")
        if any(os.path.exists(p) for p in (asl_path, xlsx_path)):
            if not messagebox.askyesno("Overwrite?", f"{base}.asl / {base}.xlsx already exist in:\n{out_dir}\n\nOverwrite them?",
                                       parent=self):
                return
        try:
            os.makedirs(out_dir, exist_ok=True)
            _, events, cfg = self._build_config()
            write_asl_file(asl_path, cfg)
            build_excel(xlsx_path, events)
        except Exception as e:  # noqa: BLE001
            messagebox.showerror("Could not write files", str(e), parent=self)
            return
        save_settings({**load_settings(), "output_dir": out_dir})
        # remember network names for next time
        net_map = load_network_map()
        for k, v in self.net_vars.items():
            net_map[k] = v.get().strip()
        save_network_map(net_map)
        messagebox.showinfo("Done", f"Config file:\n  {asl_path}\n\nExcel file:\n  {xlsx_path}", parent=self)


def main():
    if openpyxl is None:
        print("This script requires openpyxl:  pip install openpyxl")
        sys.exit(1)
    App().mainloop()


if __name__ == "__main__":
    main()
